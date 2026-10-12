"""DOCX (Word) text extraction."""
from __future__ import annotations

import io
import logging

from docx import Document

from shared.schemas import OCRResult

logger = logging.getLogger(__name__)


def extract_docx(data: bytes) -> OCRResult:
    """Extract text from a Word document."""
    doc = Document(io.BytesIO(data))

    paragraphs: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Also extract tables (often contain structured data)
    tables_text: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                tables_text.append(row_text)

    # Combine: paragraphs first, then tables
    body = "\n".join(paragraphs)
    if tables_text:
        body += "\n\n--- Tables ---\n" + "\n".join(tables_text)

    return OCRResult(
        page_count=1,  # DOCX doesn't have page numbers reliably
        full_text=body,
        words=[],      # bboxes not meaningful outside rendering
        regions=[],
    )
