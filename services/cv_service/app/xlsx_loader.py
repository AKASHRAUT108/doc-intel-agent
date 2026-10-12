"""XLSX (Excel) text extraction."""
from __future__ import annotations

import io
import logging

from openpyxl import load_workbook

from shared.schemas import OCRResult

logger = logging.getLogger(__name__)

MAX_ROWS_PER_SHEET = 500  # prevent giant sheets from blowing up memory


def extract_xlsx(data: bytes) -> OCRResult:
    """Extract text from an Excel workbook."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)

    sections: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines: list[str] = [f"### Sheet: {sheet_name}"]

        row_count = 0
        for row in ws.iter_rows(values_only=True):
            row_count += 1
            if row_count > MAX_ROWS_PER_SHEET:
                lines.append(f"... ({MAX_ROWS_PER_SHEET}+ rows, truncated)")
                break

            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))

        sections.append("\n".join(lines))

    wb.close()

    return OCRResult(
        page_count=len(wb.sheetnames),
        full_text="\n\n".join(sections),
        words=[],
        regions=[],
    )
